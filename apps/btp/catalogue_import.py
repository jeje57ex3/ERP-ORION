"""Import d'un catalogue fournisseur publié sur le web.

Le fournisseur met un fichier en ligne — CSV, Excel ou JSON — et l'ERP le
télécharge pour mettre à jour son catalogue. Un seul mécanisme sert tous les
fournisseurs : ce qui diffère d'un à l'autre, ce sont les noms de colonnes,
décrits par `CatalogueSource.colonnes`.

**La règle sur les prix, qui gouverne tout l'import.** Le prix d'ACHAT vient
du fournisseur : il est toujours mis à jour. Le prix de VENTE est celui de
l'exploitant, il porte sa marge : il n'est rempli que s'il est vide, et
jamais écrasé. Sans cette asymétrie, une synchronisation nocturne effacerait
en silence toute la tarification — et personne ne s'en apercevrait avant de
lire un devis parti chez un client.

**Rien n'est jamais supprimé.** Un article que le fournisseur ne publie plus
est signalé, pas effacé : il peut figurer dans un devis déjà envoyé.

**L'adresse est fournie par l'utilisateur et récupérée par le SERVEUR.**
C'est une requête sortante depuis une machine qui atteint la base de
données, Redis, Ollama et, sur un hébergeur, les métadonnées d'instance.
Une adresse pointant vers `127.0.0.1`, `10.0.0.0/8` ou `169.254.169.254`
transformerait cet import en lecteur du réseau interne. `_verifier_adresse`
refuse donc toute adresse qui résout vers une plage privée, et chaque
redirection est revérifiée.

Limite assumée et non masquée : entre la vérification du nom et la connexion
subsiste une fenêtre de réassociation DNS. La fermer demanderait de se
connecter à l'adresse IP validée en forçant l'en-tête Host, ce qui casse la
vérification du certificat en HTTPS. Comme seule une personne
administratrice configure une source, le compromis retenu est la
vérification de toutes les adresses résolues plus le contrôle de chaque
redirection.
"""

import csv
import io
import ipaddress
import json
import socket
from decimal import Decimal, InvalidOperation

from apps.core.nombres import nombre_francais
from urllib.parse import urlparse

TAILLE_MAX = 25 * 1024 * 1024
"""25 Mo. Un tarif de négoce complet pèse quelques mégaoctets ; au-delà, on
arrête plutôt que de remplir la mémoire du serveur."""

DELAI_CONNEXION = 10
DELAI_LECTURE = 60

REDIRECTIONS_MAX = 3

ENCODAGES_ESSAYES = ('utf-8-sig', 'utf-8', 'cp1252')
"""Dans cet ordre. Les tarifs des négoces français sortent souvent de
tableurs Windows, donc en cp1252 : essayé en dernier car il décode presque
n'importe quoi sans lever d'erreur, et masquerait un vrai UTF-8."""

NATURES_RECONNUES = {
    'materiau': 'material', 'matériau': 'material', 'material': 'material',
    'materiel': 'material', 'matériel': 'material', 'materiaux': 'material',
    'matériaux': 'material',
    'fourniture': 'supply', 'fournitures': 'supply', 'supply': 'supply',
    'consommable': 'supply',
    'action': 'action', 'pose': 'action', 'prestation': 'action',
    'main d’oeuvre': 'action', "main d'oeuvre": 'action',
    'main-d’oeuvre': 'action', "main-d'oeuvre": 'action',
    'service': 'action',
}


class CatalogueError(Exception):
    """Ce qui empêche l'import, dit dans des termes exploitables."""


# ─── L'adresse ───────────────────────────────────────────────────────────────

def _verifier_adresse(url):
    """Refuse une adresse qui mène au réseau interne. Rend le nom d'hôte.

    Toutes les adresses résolues sont contrôlées, et non la première : un
    nom peut résoudre à la fois vers une adresse publique et vers
    `127.0.0.1`, et ne vérifier que la première laisserait passer.
    """
    analysee = urlparse(url or '')
    if analysee.scheme not in ('http', 'https'):
        raise CatalogueError(
            'L’adresse doit commencer par http:// ou https://.')

    hote = analysee.hostname
    if not hote:
        raise CatalogueError('Adresse sans nom de serveur.')

    try:
        resolues = socket.getaddrinfo(hote, None)
    except socket.gaierror:
        raise CatalogueError(f'Nom de serveur introuvable : {hote}.')

    for infos in resolues:
        brute = infos[4][0]
        try:
            adresse = ipaddress.ip_address(brute)
        except ValueError:
            continue
        if (adresse.is_private or adresse.is_loopback
                or adresse.is_link_local or adresse.is_reserved
                or adresse.is_multicast or adresse.is_unspecified):
            raise CatalogueError(
                f'L’adresse {hote} mène au réseau interne ({brute}) : '
                'refusée. Un catalogue fournisseur est public.')

    return hote


def telecharger(source, session=None):
    """Les octets du fichier publié par cette source."""
    return telecharger_url(source.url, session=session)


def telecharger_url(url, session=None):
    """Les octets d'une adresse publique, ou CatalogueError.

    Extraite de `telecharger` pour que la DÉTECTION d'articles sur une page
    web passe par le même contrôle : l'adresse vient de l'utilisateur dans
    les deux cas, et c'est le serveur qui la récupère dans les deux cas.
    Deux chemins de téléchargement, c'est un chemin qui finit sans garde.

    Les redirections ne sont PAS suivies automatiquement : chaque étape est
    revérifiée, sinon une adresse publique redirigeant vers `127.0.0.1`
    contournerait tout le contrôle.
    """
    import requests

    session = session or requests.Session()
    for _ in range(REDIRECTIONS_MAX + 1):
        _verifier_adresse(url)
        try:
            reponse = session.get(
                url, timeout=(DELAI_CONNEXION, DELAI_LECTURE), stream=True,
                allow_redirects=False,
                headers={'User-Agent': 'Orion-ERP/1.0 (import catalogue)'},
            )
        except requests.RequestException as exc:
            raise CatalogueError(f'Téléchargement impossible : {exc}')

        if reponse.is_redirect or reponse.is_permanent_redirect:
            suivante = reponse.headers.get('Location')
            reponse.close()
            if not suivante:
                raise CatalogueError('Redirection sans destination.')
            url = requests.compat.urljoin(url, suivante)
            continue

        if reponse.status_code != 200:
            reponse.close()
            raise CatalogueError(
                f'Le serveur a répondu {reponse.status_code}.')

        # Lecture bornee : un fichier de 2 Go ne doit pas tenir en memoire.
        morceaux = []
        total = 0
        for morceau in reponse.iter_content(chunk_size=65536):
            total += len(morceau)
            if total > TAILLE_MAX:
                reponse.close()
                raise CatalogueError(
                    f'Fichier trop volumineux (plus de '
                    f'{TAILLE_MAX // 1024 // 1024} Mo).')
            morceaux.append(morceau)
        reponse.close()

        octets = b''.join(morceaux)
        if not octets:
            raise CatalogueError('Le fichier téléchargé est vide.')
        return octets

    raise CatalogueError('Trop de redirections.')


# ─── La lecture ──────────────────────────────────────────────────────────────

def _decoder(octets, encodage):
    if encodage and encodage != 'auto':
        try:
            return octets.decode(encodage)
        except (LookupError, UnicodeDecodeError):
            raise CatalogueError(
                f'Le fichier ne se lit pas en {encodage}.')
    for essai in ENCODAGES_ESSAYES:
        try:
            return octets.decode(essai)
        except UnicodeDecodeError:
            continue
    raise CatalogueError(
        'Encodage du fichier non reconnu. Précisez-le sur la source.')


def _lire_csv(octets, source):
    texte = _decoder(octets, source.encodage)
    delimiteur = (source.delimiteur or '').strip()
    if not delimiteur:
        # Deviner, en se limitant aux separateurs plausibles : le sniffer
        # de csv prend volontiers une lettre frequente pour un separateur.
        try:
            delimiteur = csv.Sniffer().sniff(
                texte[:4096], delimiters=';,\t|').delimiter
        except csv.Error:
            delimiteur = ';'

    lecteur = csv.DictReader(io.StringIO(texte), delimiter=delimiteur)
    if not lecteur.fieldnames:
        raise CatalogueError('Le fichier n’a pas de ligne d’en-tête.')
    return [dict(ligne) for ligne in lecteur]


def _lire_excel(octets, source):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise CatalogueError('Lecture Excel indisponible sur ce serveur.')

    try:
        classeur = load_workbook(io.BytesIO(octets), read_only=True,
                                 data_only=True)
    except Exception as exc:
        raise CatalogueError(f'Fichier Excel illisible : {exc}')

    feuille = classeur[classeur.sheetnames[0]]
    lignes = feuille.iter_rows(values_only=True)
    try:
        entetes = [str(c).strip() if c is not None else ''
                   for c in next(lignes)]
    except StopIteration:
        raise CatalogueError('Le fichier Excel est vide.')

    resultat = []
    for valeurs in lignes:
        if valeurs is None or all(v is None for v in valeurs):
            continue
        resultat.append({
            entete: valeurs[index] if index < len(valeurs) else None
            for index, entete in enumerate(entetes) if entete
        })
    classeur.close()
    return resultat


def _lire_json(octets, source):
    texte = _decoder(octets, source.encodage)
    try:
        charge = json.loads(texte)
    except ValueError as exc:
        raise CatalogueError(f'JSON illisible : {exc}')

    if isinstance(charge, list):
        lignes = charge
    elif isinstance(charge, dict):
        # Forme courante : {"produits": [...]}. On prend la seule liste
        # presente ; s'il y en a plusieurs, on ne devine pas.
        listes = [v for v in charge.values() if isinstance(v, list)]
        if len(listes) != 1:
            raise CatalogueError(
                'JSON attendu : une liste d’articles, ou un objet '
                'contenant une seule liste.')
        lignes = listes[0]
    else:
        raise CatalogueError('JSON attendu : une liste d’articles.')

    return [ligne for ligne in lignes if isinstance(ligne, dict)]


def analyser(octets, source):
    """Les lignes du fichier, sous forme de dictionnaires."""
    lecteurs = {'csv': _lire_csv, 'excel': _lire_excel, 'json': _lire_json}
    lecteur = lecteurs.get(source.format)
    if lecteur is None:
        raise CatalogueError(f'Format inconnu : {source.format}.')
    return lecteur(octets, source)


# ─── Les valeurs ─────────────────────────────────────────────────────────────

def _texte(valeur):
    if valeur is None:
        return ''
    return str(valeur).strip()


# La lecture des nombres vit desormais dans apps/core/nombres.py : elle
# sert aussi a l'extraction des factures PDF, et une deuxieme copie aurait
# diverge a la premiere correction -- l'une lisant « 1.234,56 » et l'autre
# « 1,23456 », sur des montants qui partent en comptabilite. L'alias garde
# les appels et les tests de ce module inchanges.
_nombre = nombre_francais


def _nature(valeur, defaut):
    cle = _texte(valeur).lower()
    return NATURES_RECONNUES.get(cle, defaut)


def _valeurs_de_la_ligne(ligne, source):
    """Les champs de l'ERP, lus dans la ligne selon la correspondance."""
    mapping = source.colonne_du_champ
    lues = {}
    for champ, colonne in mapping.items():
        if colonne in ligne:
            lues[champ] = ligne[colonne]
    return lues


# ─── L'import ────────────────────────────────────────────────────────────────

class Rapport:
    """Ce qu'un import a fait, avant de l'enregistrer."""

    def __init__(self):
        self.lignes_lues = 0
        self.crees = 0
        self.mis_a_jour = 0
        self.inchanges = 0
        self.ignores = 0
        self.erreurs = []

    def ignorer(self, numero, raison):
        self.ignores += 1
        # Borne : un fichier entierement mal mappe produirait des milliers
        # d'erreurs identiques, illisibles et lourdes a stocker.
        if len(self.erreurs) < 50:
            self.erreurs.append(f'Ligne {numero} : {raison}')

    @property
    def statut(self):
        if self.crees or self.mis_a_jour or self.inchanges:
            return 'partiel' if self.ignores else 'succes'
        return 'echec'


def importer(source, lignes, a_blanc=False):
    """Crée ou met à jour les articles, et rend le rapport.

    `a_blanc` lit, compare et compte sans rien écrire : c'est ce qui permet
    de vérifier une correspondance de colonnes avant de toucher à un
    catalogue de plusieurs centaines d'articles.
    """
    from django.db import transaction

    from .models import Material

    rapport = Rapport()
    mapping = source.colonne_du_champ
    if 'nom' not in mapping:
        raise CatalogueError(
            'La correspondance doit au moins indiquer la colonne de la '
            'désignation.')

    a_ecrire = []
    for numero, ligne in enumerate(lignes, start=2):  # 1 = en-tete
        rapport.lignes_lues += 1
        lues = _valeurs_de_la_ligne(ligne, source)

        nom = _texte(lues.get('nom'))
        if not nom:
            rapport.ignorer(numero, 'désignation vide')
            continue

        reference = _texte(lues.get('reference'))[:30]

        # Rapprochement : la reference si elle existe, sinon la
        # designation. La reference est plus sure -- deux articles peuvent
        # porter le meme nom -- mais tous les tarifs n'en publient pas.
        existant = None
        if reference:
            existant = Material.objects.filter(
                company=source.company, reference__iexact=reference).first()
        else:
            existant = Material.objects.filter(
                company=source.company, name__iexact=nom,
                reference='').first()

        prix_achat = _nombre(lues.get('prix_achat'))
        prix_vente = _nombre(lues.get('prix_vente'))

        champs = {
            'name': nom[:300],
            'kind': _nature(lues.get('nature'), source.nature_par_defaut),
        }
        if 'categorie' in lues:
            champs['category'] = _texte(lues.get('categorie'))[:100]
        if 'unite' in lues and _texte(lues.get('unite')):
            champs['unit'] = _texte(lues.get('unite'))[:20]
        if 'reference_fournisseur' in lues:
            champs['supplier_reference'] = _texte(
                lues.get('reference_fournisseur'))[:100]
        if 'description' in lues:
            champs['description'] = _texte(lues.get('description'))
        if prix_achat is not None:
            # Le prix d'ACHAT vient du fournisseur : toujours mis a jour.
            champs['purchase_price'] = prix_achat

        if existant is None:
            champs.update({
                'company': source.company,
                'reference': reference,
                'catalogue_source': source,
                'supplier': source.supplier,
            })
            # A la CREATION seulement, le prix de vente propose par le
            # fournisseur est repris : il n'y a pas encore de marge a
            # preserver.
            if prix_vente is not None:
                champs['sale_price'] = prix_vente
            rapport.crees += 1
            if not a_blanc:
                a_ecrire.append(('creer', Material(**champs)))
            continue

        # Le prix de VENTE n'est JAMAIS ecrase : il porte la marge de
        # l'exploitant. Il n'est rempli que s'il est encore vide.
        if prix_vente is not None and not existant.sale_price:
            champs['sale_price'] = prix_vente

        modifies = [
            nom_champ for nom_champ, valeur in champs.items()
            if getattr(existant, nom_champ) != valeur
        ]
        if not modifies:
            rapport.inchanges += 1
            continue

        rapport.mis_a_jour += 1
        if not a_blanc:
            for nom_champ, valeur in champs.items():
                setattr(existant, nom_champ, valeur)
            if existant.catalogue_source_id is None:
                existant.catalogue_source = source
            a_ecrire.append(('modifier', existant))

    if not a_blanc and a_ecrire:
        # Tout ou rien : un catalogue a moitie importe laisse des prix
        # d'une version et des prix de l'autre, sans moyen de savoir
        # lesquels.
        with transaction.atomic():
            for action, objet in a_ecrire:
                objet.save()

    return rapport


def executer(source, a_blanc=False, utilisateur=None):
    """Télécharge, lit, importe, et journalise. Rend le CatalogueImport.

    L'échec est journalisé lui aussi : une source qui ne répond plus depuis
    trois semaines doit se voir, et un import qui échoue en silence est
    pire qu'un import absent.
    """
    from django.utils import timezone

    from .models import CatalogueImport

    journal = CatalogueImport.objects.create(
        source=source, a_blanc=a_blanc, lance_par=utilisateur)

    try:
        octets = telecharger(source)
        lignes = analyser(octets, source)
        rapport = importer(source, lignes, a_blanc=a_blanc)
    except CatalogueError as exc:
        journal.statut = 'echec'
        journal.message = str(exc)[:300]
        journal.termine_le = timezone.now()
        journal.save()
        return journal

    journal.statut = rapport.statut
    journal.lignes_lues = rapport.lignes_lues
    journal.crees = rapport.crees
    journal.mis_a_jour = rapport.mis_a_jour
    journal.inchanges = rapport.inchanges
    journal.ignores = rapport.ignores
    journal.erreurs = rapport.erreurs
    journal.termine_le = timezone.now()
    journal.save()

    if not a_blanc:
        source.derniere_execution = timezone.now()
        source.save(update_fields=['derniere_execution'])

    return journal


def articles_disparus(source):
    """Les articles de cette source absents du dernier fichier publié.

    Rien n'est supprimé : un article peut figurer dans un devis déjà
    envoyé. On se contente de les nommer, pour que l'exploitant décide.
    """
    from .models import Material

    try:
        octets = telecharger(source)
        lignes = analyser(octets, source)
    except CatalogueError:
        return []

    mapping = source.colonne_du_champ
    colonne_ref = mapping.get('reference')
    colonne_nom = mapping.get('nom')

    references = {_texte(l.get(colonne_ref)).lower()
                  for l in lignes if colonne_ref}
    noms = {_texte(l.get(colonne_nom)).lower() for l in lignes if colonne_nom}

    disparus = []
    for article in Material.objects.filter(catalogue_source=source):
        cle = (article.reference or '').lower()
        if cle and cle in references:
            continue
        if not cle and article.name.lower() in noms:
            continue
        disparus.append(article)
    return disparus
